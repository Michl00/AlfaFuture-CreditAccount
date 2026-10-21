import io
import json
import os
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker

import test_credit_api as credit_tests
from services import credit_api as credit
from services.checks import run_checks
from services.db import BankStaff, User
from services.documents import parse_fields


class BankTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = credit_tests.CreditAPITests.asyncSetUp
    asyncTearDown = credit_tests.CreditAPITests.asyncTearDown
    register = credit_tests.CreditAPITests.register
    login = credit_tests.CreditAPITests.login
    prepare = credit_tests.CreditAPITests.prepare
    assess = credit_tests.CreditAPITests.assess

    async def submit(self, **changes):
        _, data = await self.prepare()
        data.update(changes)
        assessment = await self.assess(data)
        self.assertEqual(assessment.json()['status'], 'ready', assessment.text)
        response = await self.client.post('/credit/confirm', json={
            'assessment_token': assessment.json()['assessment_token'], 'data': data, 'consent': True})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    async def employee(self, email='staff@example.com'):
        await self.client.post('/auth/logout')
        await self.client.post('/auth/register', json={'email': email, 'password': 'test-password-123',
            'username': 'employee', 'inn': '7701234560'})
        sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        async with sessions() as session:
            user = await session.scalar(select(User).where(User.email == email))
            if await session.get(BankStaff, user.id) is None:
                session.add(BankStaff(user_id=user.id))
                await session.commit()
        await self.client.post('/auth/login', data={'username': email, 'password': 'test-password-123'})

    async def action(self, item, action, comment='Результаты рассмотрены сотрудником'):
        return await self.client.post(f"/bank/applications/{item['id']}/actions", json={
            'version': item['version'], 'action': action, 'comment': comment})

    async def test_role_is_required_and_cannot_be_assigned_through_profile(self):
        self.assertEqual((await self.client.get('/bank/me')).status_code, 401)
        await self.login()
        self.assertEqual((await self.client.get('/bank/applications')).status_code, 403)
        self.assertEqual((await self.client.patch('/users/me', json={'role': 'bank_employee'})).status_code, 422)
        self.assertEqual((await self.client.get('/bank/me')).status_code, 403)

    async def test_console_can_grant_and_revoke_staff_access(self):
        import manage
        await self.login()
        sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        with patch.object(manage, 'async_session_maker', sessions), patch.object(manage, 'create_db_and_tables', new_callable=AsyncMock):
            await manage.set_staff('demo@example.com')
            await manage.set_staff('demo@example.com')
            self.assertEqual((await self.client.get('/bank/me')).status_code, 200)
            await manage.set_staff('demo@example.com', revoke=True)
            self.assertEqual((await self.client.get('/bank/me')).status_code, 403)

    async def test_queue_manual_filter_claim_decision_and_final_immutability(self):
        item = await self.submit(inn='7701234578')
        own = (await self.client.get('/credit/applications')).json()
        self.assertEqual(own[0]['id'], item['id'])
        self.assertTrue(item['score']['manual_review'])
        await self.employee()
        queue = (await self.client.get('/bank/applications?manual_only=true')).json()
        self.assertEqual(queue['items'][0]['id'], item['id'])
        self.assertEqual((await self.action(item, 'approve')).status_code, 409)
        claimed = (await self.action(item, 'claim')).json()
        self.assertEqual(claimed['status'], 'in_review')
        self.assertEqual(claimed['applicant']['email'], 'demo@example.com')
        self.assertEqual((await self.action(item, 'approve')).status_code, 409)
        self.assertEqual((await self.action(claimed, 'approve', ' ')).status_code, 422)
        approved = (await self.action(claimed, 'approve')).json()
        self.assertEqual(approved['status'], 'approved')
        self.assertEqual([event['action'] for event in approved['events']], ['submit', 'claim', 'approve'])
        self.assertEqual((await self.action(approved, 'reject')).status_code, 409)

    async def test_other_employee_cannot_claim_an_assigned_application(self):
        item = await self.submit()
        await self.employee()
        claimed = (await self.action(item, 'claim')).json()
        await self.employee('second@example.com')
        self.assertEqual((await self.action(claimed, 'claim')).status_code, 409)
        self.assertEqual((await self.action(claimed, 'reject')).status_code, 409)

    async def test_request_info_client_reply_requeues_and_keeps_audit(self):
        item = await self.submit()
        await self.employee()
        claimed = (await self.action(item, 'claim')).json()
        requested = (await self.action(claimed, 'request_info', 'Уточните источники выручки')).json()
        self.assertEqual(requested['status'], 'needs_info')
        await self.client.post('/auth/logout')
        await self.client.post('/auth/login', data={'username': 'demo@example.com', 'password': 'test-password-123'})
        payload = {'version': requested['version'], 'comment': 'Выручка подтверждается договорами поставки', 'document_ids': []}
        reply = await self.client.post(f"/credit/applications/{item['id']}/clarification", json=payload)
        self.assertEqual(reply.status_code, 200, reply.text)
        self.assertEqual(reply.json()['status'], 'bank_review')
        self.assertIsNone(reply.json()['assigned_to'])
        self.assertEqual(reply.json()['events'][-1]['action'], 'clarification')
        self.assertEqual((await self.client.post(f"/credit/applications/{item['id']}/clarification", json=payload)).status_code, 409)

    async def test_documents_are_private_and_staff_can_read_only_attached_files(self):
        _, data = await self.prepare()
        document = json.dumps({'company': 'Компания из документа', 'inn': '7701234560', 'revenue': 1300000})
        response = await self.client.post('/credit/documents', files={'file': ('company.json', document, 'application/json')})
        self.assertEqual(response.status_code, 201, response.text)
        doc = response.json()
        self.assertEqual(doc['fields']['revenue'], 1300000)
        path = '/credit/documents/' + doc['id']
        self.assertEqual((await self.client.get(path)).status_code, 200)
        await self.employee()
        self.assertEqual((await self.client.get(path)).status_code, 404)
        await self.client.post('/auth/logout')
        await self.client.post('/auth/login', data={'username': 'demo@example.com', 'password': 'test-password-123'})
        data['document_ids'] = [doc['id']]
        assessed = (await self.assess(data)).json()
        self.assertTrue(assessed['score']['manual_review'])
        self.assertEqual(assessed['checks']['documents']['data']['conflicts'][0]['field'], 'revenue')
        confirmed = await self.client.post('/credit/confirm', json={'data': data,
            'assessment_token': assessed['assessment_token'], 'consent': True})
        item = confirmed.json()
        await self.employee('staff@example.com')
        downloaded = await self.client.get(path)
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.headers['x-content-type-options'], 'nosniff')
        self.assertEqual((await self.client.get('/credit/applications/' + item['id'])).status_code, 404)
        detail = (await self.client.get('/bank/applications/' + item['id'])).json()
        self.assertEqual(detail['documents'][0]['id'], doc['id'])
        # An ordinary unrelated client cannot read either the file or the application.
        await self.client.post('/auth/logout')
        await self.client.post('/auth/register', json={'email': 'other@example.com', 'password': 'test-password-123', 'username': 'other', 'inn': '7701234560'})
        await self.client.post('/auth/login', data={'username': 'other@example.com', 'password': 'test-password-123'})
        self.assertEqual((await self.client.get(path)).status_code, 404)
        self.assertEqual((await self.client.get('/credit/applications/' + item['id'])).status_code, 404)
        self.assertEqual((await self.client.get('/credit/applications')).json(), [])

    async def test_upload_rejects_unsupported_empty_and_unreadable_files(self):
        await self.login()
        for filename, content, expected in [('a.exe', b'abc', 415), ('a.txt', b'', 413),
                ('a.json', b'{}', 422), ('a.pdf', b'not a PDF', 422), ('a.txt', b'x' * (5 * 1024 * 1024 + 1), 413)]:
            result = await self.client.post('/credit/documents', files={'file': (filename, content)})
            self.assertEqual(result.status_code, expected, result.text)

    async def test_inactive_registry_blocks_without_model_call(self):
        _, data = await self.prepare()
        data['inn'] = '7701234592'
        with patch.object(credit, 'ask_model', new_callable=AsyncMock) as ask:
            result = await self.client.post('/credit/assessments', json=data)
        self.assertEqual(result.json()['status'], 'blocked')
        self.assertTrue(result.json()['score']['hard_stop'])
        ask.assert_not_called()

    async def test_model_receives_mock_checks_without_company_identity(self):
        _, data = await self.prepare()
        data['inn'] = '7701234585'
        with patch.object(credit, 'ask_model', new_callable=AsyncMock, return_value=credit_tests.READY) as ask:
            result = await self.client.post('/credit/assessments', json=data)
        self.assertEqual(result.status_code, 200, result.text)
        sent = ask.call_args.args[1]
        self.assertTrue(sent['external_checks']['mock'])
        self.assertTrue(sent['scoring']['manual_review'])
        self.assertNotIn('inn', sent['company_financials'])
        self.assertNotIn('company', sent['company_financials'])


class DocumentParserTests(unittest.TestCase):
    def test_pdf_with_text_layer(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 20 200 Td (inn: 7701234560) Tj 0 -20 Td (revenue: 1200000) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)
        fields, _ = parse_fields(output.getvalue(), '.pdf')
        self.assertEqual(fields['inn'], '7701234560')
        self.assertEqual(fields['revenue'], 1200000)

    def test_json_csv_and_text_fields(self):
        for content, extension in [(b'{"revenue": 1.2e6, "age":24}', '.json'),
                ('Выручка;Возраст бизнеса\n1 200 000;24\n'.encode(), '.csv'),
                ('Выручка: 1 200 000\nВозраст бизнеса: 24'.encode(), '.txt')]:
            fields, _ = parse_fields(content, extension)
            self.assertEqual(fields['revenue'], 1200000)
            self.assertEqual(fields['age'], 24)

    def test_docx_and_xlsx_tables(self):
        from docx import Document
        from openpyxl import Workbook
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [('ИНН', '7701234560'), ('Выручка', '1 200 000')]):
            for cell, value in zip(row.cells, values): cell.text = value
        book = Workbook()
        book.active.append(['ИНН', '7701234560'])
        book.active.append(['Выручка', 1200000])
        for obj, extension in [(doc, '.docx'), (book, '.xlsx')]:
            stream = io.BytesIO()
            obj.save(stream)
            fields, _ = parse_fields(stream.getvalue(), extension)
            self.assertEqual(fields['inn'], '7701234560')
            self.assertEqual(fields['revenue'], 1200000)

    def test_mock_scenarios_and_document_conflicts_affect_scoring(self):
        data = {'inn': '7701234560', 'revenue': 1200000, 'expenses': 700000, 'debt': 20000, 'age': 24}
        terms = {'amount': 1500000, 'rate': 24.9, 'months': 24}
        clean_checks, clean = run_checks(data, terms)
        self.assertFalse(clean['manual_review'])
        for inn in ('7701234578', '7701234585'):
            checks, score = run_checks({**data, 'inn': inn}, terms)
            self.assertTrue(score['manual_review'])
            self.assertLess(score['value'], clean['value'])
        _, stopped = run_checks({**data, 'inn': '7701234592'}, terms)
        self.assertTrue(stopped['hard_stop'])
        _, mismatch = run_checks(data, terms, [SimpleNamespace(id='doc', fields={'revenue': 100000})])
        self.assertTrue(mismatch['manual_review'])
        self.assertLess(mismatch['value'], clean['value'])

    def test_comet_model_default_and_key_alias(self):
        from config.config import load_config
        with patch.dict(os.environ, {'SECRET': 'test-secret', 'COOKIE_SECURE': 'false', 'COMETAPI_KEY': 'new-test-key', 'COMETAPI_API_KEY': 'old-test-key'}, clear=True), patch('environs.Env.read_env'):
            configured = load_config()
            self.assertEqual(configured.comet.model, 'qwen3.8-flash')
            self.assertEqual(configured.comet.api_key, 'new-test-key')
            os.environ['COMETAPI_KEY'] = ''
            self.assertEqual(load_config().comet.api_key, 'old-test-key')
            self.assertEqual(load_config().comet.timeout_seconds, 180)
            os.environ['COMETAPI_TIMEOUT_SECONDS'] = '240'
            self.assertEqual(load_config().comet.timeout_seconds, 240)
            for invalid in ('0', '301', 'nan'):
                os.environ['COMETAPI_TIMEOUT_SECONDS'] = invalid
                with self.assertRaises(ValueError):
                    load_config()
