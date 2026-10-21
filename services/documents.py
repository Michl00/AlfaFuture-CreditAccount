"""Bounded text/table extraction, followed by explicit client confirmation."""
import csv
import io
import json
import re
import time
import zipfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import PurePath
from typing import Annotated
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from services.db import BankStaff, CreditApplication, UploadedDocument, User, get_async_session
from services.services import valid_inn
from services.users import current_active_user

router = APIRouter(prefix='/credit/documents', tags=['documents'])
Session = Annotated[AsyncSession, Depends(get_async_session)]
Owner = Annotated[User, Depends(current_active_user)]
MAX_SIZE = 5 * 1024 * 1024
ALIASES = {
    'company': ('company', 'компания', 'организация', 'наименование', 'название компании'),
    'inn': ('inn', 'инн'),
    'revenue': ('revenue', 'выручка', 'среднемесячная выручка', 'доходы'),
    'expenses': ('expenses', 'расходы', 'среднемесячные расходы'),
    'debt': ('debt', 'кредитные платежи', 'платежи по кредитам', 'действующие платежи'),
    'age': ('age', 'возраст бизнеса', 'возраст бизнеса в месяцах', 'месяцев работы'),
}


def zip_guard(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 500 or sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError('Слишком большой распакованный документ')
    except zipfile.BadZipFile as error:
        raise ValueError('Повреждённый документ') from error


def extract_text(content, extension):
    if extension in ('.json', '.txt', '.csv'):
        try: return content.decode('utf-8-sig')
        except UnicodeDecodeError: return content.decode('cp1251')
    if extension == '.pdf':
        from pypdf import PdfReader
        if not content.startswith(b'%PDF-'):
            raise ValueError('Повреждённый PDF')
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 20:
            raise ValueError('Нужен незашифрованный PDF не более 20 страниц')
        return '\n'.join((page.extract_text() or '')[:50000] for page in reader.pages)[:200000]
    zip_guard(content)
    if extension == '.docx':
        from docx import Document
        document = Document(io.BytesIO(content))
        lines = [item.text for item in document.paragraphs]
        lines.extend(' : '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        return '\n'.join(lines)[:200000]
    if extension == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return '\n'.join(' : '.join(str(value) for value in row if value is not None)
                for sheet in book.worksheets[:5] for row in sheet.iter_rows(max_row=1000, max_col=10, values_only=True))[:200000]
        finally: book.close()
    raise ValueError('Поддерживаются PDF, DOCX, XLSX, CSV, TXT и JSON')


def parse_fields(content, extension):
    text = extract_text(content, extension)
    raw = {}
    if extension == '.json':
        raw = json.loads(text)
        if not isinstance(raw, dict): raise ValueError('JSON должен содержать объект с полями')
    elif extension == '.csv':
        delimiter = ';' if text.count(';') > text.count(',') else ','
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))[:1000]
        aliases = {alias for values in ALIASES.values() for alias in values}
        if len(rows) == 2 and all(key.strip().lower() in aliases for key in rows[0]):
            raw = dict(zip(rows[0], rows[1]))
        else: raw = {row[0].strip(): row[1].strip() for row in rows if len(row) >= 2}
    else:
        labels = '|'.join(re.escape(alias) for aliases in ALIASES.values() for alias in aliases)
        pattern = rf'^\s*({labels})\s*(?::|=|\t|\|)?\s*(.*?)\s*$'
        lines = text.splitlines()
        for index, line in enumerate(lines):
            match = re.match(pattern, line, re.I)
            if match:
                value = match.group(2).strip()
                if not value and index + 1 < len(lines): value = lines[index + 1].strip()
                raw[match.group(1).lower()] = value
    normalized = {str(key).strip().lower(): value for key, value in raw.items()}
    fields, warnings = {}, []
    for name, aliases in ALIASES.items():
        value = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if value is None or isinstance(value, (dict, list, bool)): continue
        if name == 'company':
            value = str(value).strip()
            if 2 <= len(value) <= 120: fields[name] = value
            else: warnings.append('Проверьте название компании')
        elif name == 'inn':
            value = re.sub(r'\s', '', str(value))
            if valid_inn(value): fields[name] = value
            else: warnings.append('ИНН из документа не прошёл проверку')
        else:
            try:
                if isinstance(value, (int, float)):
                    numeric = Decimal(str(value))
                else:
                    number = re.search(r'-?\d[\d\s\u00a0]*(?:[.,]\d+)?', str(value))
                    if not number: raise ValueError()
                    numeric = Decimal(re.sub(r'\s', '', number.group()).replace(',', '.'))
                integer = int(numeric.quantize(Decimal(1), rounding=ROUND_HALF_UP))
                maximum = 1200 if name == 'age' else 1000000000
                if not 0 <= integer <= maximum: raise ValueError()
                fields[name] = integer
                if numeric != integer: warnings.append(f'{name}: значение округлено до целого')
            except (ValueError, InvalidOperation): warnings.append(f'Не удалось распознать поле {name}')
    if not fields:
        raise ValueError('Поля не найдены. Используйте текстовый документ с подписями полей; сканы без текстового слоя не распознаются.')
    return fields, warnings


async def owned_documents(session, user_id, ids):
    ids = list(dict.fromkeys(ids))
    if not ids: return []
    rows = list((await session.scalars(select(UploadedDocument).where(UploadedDocument.id.in_(ids), UploadedDocument.user_id == user_id))).all())
    if len(rows) != len(ids): raise HTTPException(404, 'DOCUMENT_NOT_FOUND')
    return rows


@router.post('', status_code=201)
async def upload_document(user: Owner, session: Session, file: UploadFile = File(...)):
    content = await file.read(MAX_SIZE + 1)
    await file.close()
    if not content or len(content) > MAX_SIZE: raise HTTPException(413, 'DOCUMENT_SIZE_LIMIT')
    filename = PurePath((file.filename or 'document').replace('\\', '/')).name[:180]
    extension = PurePath(filename).suffix.lower()
    if extension not in {'.pdf', '.docx', '.xlsx', '.csv', '.txt', '.json'}: raise HTTPException(415, 'DOCUMENT_FORMAT_UNSUPPORTED')
    try:
        # CPU/file parsing runs outside FastAPI's event loop.
        import asyncio
        fields, warnings = await asyncio.to_thread(parse_fields, content, extension)
    except Exception:
        # Library errors may contain document content: return a bounded generic message.
        raise HTTPException(422, 'DOCUMENT_TEXT_NOT_RECOGNIZED') from None
    document = UploadedDocument(id=str(uuid4()), user_id=user.id, filename=filename, content=content,
        fields=fields, warnings=warnings, created_at=int(time.time() * 1000))
    session.add(document)
    await session.commit()
    return {'id': document.id, 'filename': filename, 'fields': fields, 'warnings': warnings, 'size': len(content)}


@router.get('/{document_id}')
async def download_document(document_id: str, user: Owner, session: Session):
    document = await session.get(UploadedDocument, document_id)
    if document is None: raise HTTPException(404, 'DOCUMENT_NOT_FOUND')
    if document.user_id != user.id:
        staff = await session.get(BankStaff, user.id)
        linked = False
        if staff:
            apps = (await session.scalars(select(CreditApplication).where(CreditApplication.user_id == document.user_id))).all()
            linked = any(document_id in item.payload.get('document_ids', []) for item in apps)
        if not linked: raise HTTPException(404, 'DOCUMENT_NOT_FOUND')
    return Response(document.content, media_type='application/octet-stream', headers={
        'Content-Disposition': "attachment; filename*=UTF-8''" + quote(document.filename, safe=''),
        'X-Content-Type-Options': 'nosniff', 'Cache-Control': 'no-store',
    })
